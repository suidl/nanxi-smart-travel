from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "参赛申报书-v2.md"
OUTPUT = ROOT / "楠溪智游-参赛申报书-v3.docx"
SKILL_SCRIPTS = Path(r"C:\Users\PC\.codex\plugins\cache\openai-primary-runtime\documents\26.909.22227\skills\documents\scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(83, 94, 105)
FONT = "Microsoft YaHei"


def set_font(style, size: float, color=RGBColor(26, 37, 48), bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def add_page_number(paragraph):
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    paragraph.add_run(" 页")


def set_cell_fill(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    trpr.append(flag)


def add_numbering(doc, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    aid = max(abstract_ids + [0]) + 1
    nid = max(num_ids + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(fmt)
    txt = OxmlElement("w:lvlText")
    txt.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(txt)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nid))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(aid))
    num.append(abs_ref)
    numbering.append(num)
    return nid


def set_num(paragraph, nid: int):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(nid))
    numpr.extend([ilvl, numid])
    ppr.append(numpr)


def add_table(doc, rows: list[list[str]], width: int):
    cols = len(rows[0])
    weights = {2: [2.0, 5.0], 3: [1.4, 2.6, 3.0], 4: [1.4, 1.3, 2.1, 2.2]}.get(cols, [1] * cols)
    widths = column_widths_from_weights(weights, width)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, values in enumerate(rows):
        for j, value in enumerate(values):
            cell = table.cell(i, j)
            cell.text = clean_inline(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.style = doc.styles["Table Body"]
                if i == 0:
                    for run in p.runs:
                        run.bold = True
            if i == 0:
                set_cell_fill(cell, "F4F6F9")
    set_repeat_header(table.rows[0])
    apply_table_geometry(table, widths, table_width_dxa=width, indent_dxa=120,
                         cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120})
    doc.add_paragraph(style="After Table")


def clean_inline(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1（\2）", value)
    return value.replace("`", "")


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.1)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)
section.header_distance = Cm(1.25)
section.footer_distance = Cm(1.25)
width = section_content_width_dxa(section)

normal = doc.styles["Normal"]
set_font(normal, 10.5)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, color, before, after in [
    ("Heading 1", 16, RGBColor(0, 0, 0), 16, 8),
    ("Heading 2", 13, RGBColor(0, 0, 0), 12, 6),
    ("Heading 3", 12, RGBColor(0, 0, 0), 8, 4),
]:
    style = doc.styles[name]
    set_font(style, size, color, True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.keep_with_next = True

for name, size, after in [("Table Body", 9, 0), ("After Table", 2, 4)]:
    style = doc.styles.add_style(name, 1)
    set_font(style, size)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15

header = section.header.paragraphs[0]
header.text = "2026 首届“永嘉农商杯”AI＋OPC 创新创业大赛  |  楠溪智游"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_number(footer)
for run in footer.runs:
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

lines = SOURCE.read_text(encoding="utf-8").splitlines()
title = lines[0].removeprefix("# ")
subtitle = lines[2]
meta = lines[4]

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
p.paragraph_format.space_after = Pt(14)
r = p.add_run("2026 首届“永嘉农商杯”AI＋OPC 创新创业大赛")
r.font.name = FONT
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = MUTED

p = doc.add_paragraph()
p.style = doc.styles["Title"]
set_font(doc.styles["Title"], 24, RGBColor(0, 0, 0), True)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run(title)
r.font.name = FONT
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
r = p.add_run(subtitle)
r.font.name = FONT
r.font.size = Pt(11)
r.font.color.rgb = MUTED

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(32)
r = p.add_run(meta)
r.font.name = FONT
r.font.size = Pt(9)
r.font.color.rgb = MUTED

bullet_id = add_numbering(doc, "bullet")
number_id = add_numbering(doc, "decimal")
i = 5
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith("## "):
        doc.add_paragraph(line[3:], style="Heading 1")
        i += 1
        continue
    if line.startswith("| "):
        group = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            values = [x.strip() for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", x) for x in values):
                group.append(values)
            i += 1
        if group:
            add_table(doc, group, width)
        continue
    if line.startswith("- "):
        p = doc.add_paragraph(clean_inline(line[2:]))
        set_num(p, bullet_id)
        i += 1
        continue
    match = re.match(r"^\d+\.\s+", line)
    if match:
        p = doc.add_paragraph(clean_inline(line[match.end():]))
        set_num(p, number_id)
        i += 1
        continue
    doc.add_paragraph(clean_inline(line))
    i += 1

doc.core_properties.title = "楠溪智游 楠溪江动态行程决策智能体参赛申报书"
doc.core_properties.subject = "2026 首届永嘉农商杯 AI＋OPC 创新创业大赛"
doc.core_properties.author = ""
doc.save(OUTPUT)
print(OUTPUT)
